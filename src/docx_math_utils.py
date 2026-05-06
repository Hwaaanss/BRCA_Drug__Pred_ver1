"""Helpers for writing Word-safe math-like text with python-docx.

The manuscript text contains many inline formulas written with Unicode
subscripts/superscripts (IC₅₀, 10⁻³⁰), Greek symbols, and TeX-like markers
(ŷ_d^patient).  If those are inserted as a single plain run, Word or downstream
DOCX-to-PDF converters can render them as boxes or misplaced glyphs depending on
the default font.  These helpers split math-like text into runs and apply Word's
native subscript/superscript flags plus Cambria Math for math symbols.
"""
from __future__ import annotations

from docx.oxml.ns import qn


MATH_FONT = "Cambria Math"

SUBSCRIPT_MAP = {
    "₀": "0", "₁": "1", "₂": "2", "₃": "3", "₄": "4",
    "₅": "5", "₆": "6", "₇": "7", "₈": "8", "₉": "9",
    "₊": "+", "₋": "-", "₌": "=", "₍": "(", "₎": ")",
    "ₐ": "a", "ₑ": "e", "ₕ": "h", "ᵢ": "i", "ⱼ": "j",
    "ₖ": "k", "ₗ": "l", "ₘ": "m", "ₙ": "n", "ₒ": "o",
    "ₚ": "p", "ᵣ": "r", "ₛ": "s", "ₜ": "t", "ᵤ": "u",
    "ᵥ": "v", "ₓ": "x",
}

SUPERSCRIPT_MAP = {
    "⁰": "0", "¹": "1", "²": "2", "³": "3", "⁴": "4",
    "⁵": "5", "⁶": "6", "⁷": "7", "⁸": "8", "⁹": "9",
    "⁺": "+", "⁻": "-", "⁼": "=", "⁽": "(", "⁾": ")",
    "ᵃ": "a", "ᵇ": "b", "ᶜ": "c", "ᵈ": "d", "ᵉ": "e",
    "ᶠ": "f", "ᵍ": "g", "ʰ": "h", "ⁱ": "i", "ʲ": "j",
    "ᵏ": "k", "ˡ": "l", "ᵐ": "m", "ⁿ": "n", "ᵒ": "o",
    "ᵖ": "p", "ʳ": "r", "ˢ": "s", "ᵗ": "t", "ᵘ": "u",
    "ᵛ": "v", "ʷ": "w", "ˣ": "x", "ʸ": "y", "ᶻ": "z",
}

SUBSCRIPT = str.maketrans(SUBSCRIPT_MAP)
SUPERSCRIPT = str.maketrans(SUPERSCRIPT_MAP)
SUBSCRIPT_CHARS = set(SUBSCRIPT_MAP.keys())
SUPERSCRIPT_CHARS = set(SUPERSCRIPT_MAP.keys())
MATH_CHARS = set(
    "±≤≥≈≠×÷·−→←↔∞∈∉∋∑Σ∏Π√∂∆ΔδλβαρσμθηκωΩℝ⊙⊕⊗∪∩∧∨∫"
)


def _set_run_font(run, font_name: str = MATH_FONT) -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        from docx.oxml import OxmlElement

        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), font_name)


def _clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _read_script_token(text: str, start: int) -> tuple[str, int]:
    if start >= len(text):
        return "", start
    if text[start] == "{":
        end = text.find("}", start + 1)
        if end != -1:
            return text[start + 1:end], end + 1
    end = start
    while end < len(text) and (text[end].isalnum() or text[end] in "_+-"):
        end += 1
    if end == start:
        return text[start], start + 1
    return text[start:end], end


def _looks_formula_like(text: str) -> bool:
    return any(tok in text for tok in ("=", "∈", "Σ", "||", "argmin", "where", "Loss", "L_"))


def append_math_text(paragraph, text: str, parse_ascii_scripts: bool | None = None) -> None:
    if parse_ascii_scripts is None:
        parse_ascii_scripts = _looks_formula_like(text)

    buf: list[str] = []

    def flush() -> None:
        if buf:
            paragraph.add_run("".join(buf))
            buf.clear()

    i = 0
    while i < len(text):
        ch = text[i]
        if ch in SUBSCRIPT_CHARS:
            flush()
            run = paragraph.add_run(ch.translate(SUBSCRIPT))
            run.font.subscript = True
            _set_run_font(run)
            i += 1
        elif ch in SUPERSCRIPT_CHARS:
            flush()
            run = paragraph.add_run(ch.translate(SUPERSCRIPT))
            run.font.superscript = True
            _set_run_font(run)
            i += 1
        elif parse_ascii_scripts and ch in {"_", "^"}:
            token, nxt = _read_script_token(text, i + 1)
            if token:
                flush()
                run = paragraph.add_run(token)
                if ch == "_":
                    run.font.subscript = True
                else:
                    run.font.superscript = True
                _set_run_font(run)
                i = nxt
            else:
                buf.append(ch)
                i += 1
        elif ch in MATH_CHARS:
            flush()
            run = paragraph.add_run(ch)
            _set_run_font(run)
            i += 1
        else:
            buf.append(ch)
            i += 1
    flush()


def set_paragraph_math_text(paragraph, text: str, style=None) -> None:
    if style is not None:
        try:
            paragraph.style = style
        except Exception:
            pass
    _clear_paragraph(paragraph)
    append_math_text(paragraph, text)


def add_math_paragraph(doc, text: str, style=None):
    paragraph = doc.add_paragraph(style=style)
    append_math_text(paragraph, text)
    return paragraph


def insert_math_paragraph_after(anchor_paragraph, text: str, style=None):
    paragraph = anchor_paragraph._parent.add_paragraph(style=style)
    append_math_text(paragraph, text)
    parent = anchor_paragraph._p.getparent()
    parent.remove(paragraph._p)
    idx = list(parent).index(anchor_paragraph._p)
    parent.insert(idx + 1, paragraph._p)
    return paragraph


def insert_math_paragraphs_after(anchor_paragraph, texts, style=None):
    if isinstance(texts, str):
        texts = [texts]
    inserted = []
    parent = anchor_paragraph._p.getparent()
    idx = list(parent).index(anchor_paragraph._p)
    for offset, text in enumerate(texts, start=1):
        paragraph = anchor_paragraph._parent.add_paragraph(style=style)
        append_math_text(paragraph, text)
        parent.remove(paragraph._p)
        parent.insert(idx + offset, paragraph._p)
        inserted.append(paragraph)
    return inserted


def insert_math_paragraphs_before(anchor_paragraph, texts, style=None):
    if isinstance(texts, str):
        texts = [texts]
    inserted = []
    for text in texts:
        paragraph = anchor_paragraph.insert_paragraph_before()
        set_paragraph_math_text(paragraph, text, style=style)
        inserted.append(paragraph)
    return inserted


def paragraph_needs_math_repair(paragraph) -> bool:
    text = paragraph.text
    if not text:
        return False
    if any(node.tag in {qn("w:drawing"), qn("w:pict"), qn("w:object")} for node in paragraph._p.iter()):
        return False
    if any(ch in text for ch in SUBSCRIPT_CHARS | SUPERSCRIPT_CHARS | MATH_CHARS):
        return True
    return _looks_formula_like(text) and any(marker in text for marker in ("_", "^"))


def repair_document_math_text(doc) -> int:
    """Rewrite math-like paragraphs/cells into Word-safe rich runs.

    Returns the number of paragraphs repaired.
    """
    count = 0

    def repair_paragraphs(paragraphs):
        nonlocal count
        for paragraph in paragraphs:
            if paragraph_needs_math_repair(paragraph):
                text = paragraph.text
                set_paragraph_math_text(paragraph, text)
                count += 1

    repair_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                repair_paragraphs(cell.paragraphs)
    return count
